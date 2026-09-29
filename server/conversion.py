# @index-begin
# @symbol variable/parameter: MAX_BYTES L98
# @symbol variable/parameter: MAX_TEXT L99
# @symbol variable/parameter: SUPPORTED L100
# @symbol variable/parameter: SAFE_TAGS L117
# @symbol function/class: clean_html L143
# @symbol variable/parameter: value L143
# @symbol function/class: archive_guard L154
# @symbol variable/parameter: data L154
# @symbol variable/parameter: required L154
# @symbol variable/parameter: archive L156
# @symbol variable/parameter: items L157
# @symbol function/class: image_block L176
# @symbol variable/parameter: label L176
# @symbol variable/parameter: image L178
# @symbol variable/parameter: output L182
# @symbol variable/parameter: rows L192
# @symbol function/class: table_block L192
# @symbol variable/parameter: cell L198
# @symbol function/class: extract_docx L202
# @symbol variable/parameter: document L205
# @symbol variable/parameter: blocks L206
# @symbol variable/parameter: warnings L207
# @symbol variable/parameter: item L210
# @symbol variable/parameter: equation L223
# @symbol variable/parameter: drawing L232
# @symbol variable/parameter: relation L233
# @symbol function/class: extract_pdf L245
# @symbol variable/parameter: reader L249
# @symbol variable/parameter: pdf L256
# @symbol variable/parameter: number L257
# @symbol variable/parameter: page L257
# @symbol variable/parameter: content L258
# @symbol variable/parameter: figure L268
# @symbol function/class: extract L273
# @symbol variable/parameter: name L273
# @symbol variable/parameter: extension L275
# @symbol variable/parameter: workbook L290
# @symbol variable/parameter: sheet L297
# @symbol variable/parameter: text L308
# @symbol variable/parameter: _ L319
# @symbol variable/parameter: key L322
# @symbol variable/parameter: keys L322
# @symbol variable/parameter: root L326
# @symbol variable/parameter: safe L331
# @symbol variable/parameter: part L337
# @symbol function/class: render_document L340
# @symbol variable/parameter: title L341
# @symbol variable/parameter: measurements L343
# @symbol variable/parameter: provenance L344
# @symbol variable/parameter: config L345
# @symbol variable/parameter: parts L349
# @symbol variable/parameter: section L350
# @symbol variable/parameter: warning L356
# @symbol variable/parameter: kind L359
# @symbol variable/parameter: tag L364
# @symbol variable/parameter: style L394
# @symbol function/class: ingest L398
# @symbol variable/parameter: block L404
# @symbol variable/parameter: index L404
# @symbol variable/parameter: row L406
# @symbol variable/parameter: column L407
# @symbol variable/parameter: records L408
# @symbol variable/parameter: record L409
# @symbol variable/parameter: soup L412
# @symbol variable/parameter: node L413
# @symbol variable/parameter: normalized L414
# @symbol variable/parameter: html L428
# @index-end
"""Bounded import adapters and canonical HTML generation. Symbols: docs/code-index.md.

Preserve tables, figures and equation source representations; uncertain semantics require review.
"""

import base64
import csv
import hashlib
import io
import json
import zipfile
from html import escape
from pathlib import PurePath

import nh3
import pdfplumber
from bs4 import BeautifulSoup
from defusedxml import ElementTree
from docx import Document
from lxml import etree
from markdown_it import MarkdownIt
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

from .models import WorkspaceConfig
from .units import normalize

MAX_BYTES = 4 * 1024 * 1024
MAX_TEXT = 200_000
SUPPORTED = [
    "txt",
    "md",
    "html",
    "htm",
    "csv",
    "tsv",
    "json",
    "xml",
    "docx",
    "xlsx",
    "pdf",
    "png",
    "jpg",
    "jpeg",
    "webp",
]
SAFE_TAGS = {
    "p",
    "h1",
    "h2",
    "h3",
    "h4",
    "ul",
    "ol",
    "li",
    "strong",
    "em",
    "code",
    "pre",
    "blockquote",
    "table",
    "thead",
    "tbody",
    "tr",
    "th",
    "td",
    "br",
    "hr",
    "a",
}


def clean_html(value: str) -> str:
    """Strip active content, remote assets, styles, and unsafe URL schemes."""
    return nh3.clean(
        value,
        tags=SAFE_TAGS,
        attributes={"a": {"href", "title"}},
        url_schemes={"https"},
        clean_content_tags={"script", "style", "iframe", "object"},
    )


def archive_guard(data: bytes, required: str) -> None:
    """Validate Office archive members, expansion budgets, and expected content signatures."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        items = archive.infolist()
        if len(items) > 300 or sum(item.file_size for item in items) > 16 * 1024 * 1024:
            raise ValueError("Office archive exceeds expansion limits")
        for item in items:
            if (
                item.flag_bits & 1
                or item.file_size > 8 * 1024 * 1024
                or ".." in PurePath(item.filename).parts
                or item.filename.startswith("/")
            ):
                raise ValueError("Unsafe Office archive member")
            if item.file_size / max(item.compress_size, 1) > 200:
                raise ValueError("Suspicious archive compression ratio")
            if item.filename.endswith((".xml", ".rels")):
                ElementTree.fromstring(archive.read(item))
        if required not in archive.namelist():
            raise ValueError("File signature does not match its extension")


def image_block(data: bytes, label: str) -> dict:
    """Re-encode a bounded raster to PNG; never trust uploaded metadata or SVG scripting."""
    with Image.open(io.BytesIO(data)) as image:
        if image.width * image.height > 8_000_000:
            raise ValueError("Image exceeds pixel limit")
        image.thumbnail((1400, 1400))
        output = io.BytesIO()
        image.convert("RGB").save(output, "PNG")
    return {
        "kind": "figure",
        "text": label,
        "image": base64.b64encode(output.getvalue()).decode(),
        "review": "Image/graph meaning has not been verified; request vision analysis and review.",
    }


def table_block(rows: list[list]) -> dict:
    """Bound spreadsheet/table dimensions and stringify cells without executing formulas."""
    if len(rows) > 2000 or any(len(row) > 40 for row in rows):
        raise ValueError("Table exceeds 2000 rows or 40 columns")
    return {
        "kind": "table",
        "rows": [["" if cell is None else str(cell) for cell in row] for row in rows],
    }


def extract_docx(data: bytes) -> tuple[list[dict], list[str]]:
    """Retain document order, table cells, inline images, and raw Office Math XML."""
    archive_guard(data, "word/document.xml")
    document = Document(io.BytesIO(data))
    blocks: list[dict] = []
    warnings = [
        "DOCX headers, footers, tracked changes, nested objects, and complex layout may need source review."
    ]
    for item in document.iter_inner_content():
        if hasattr(item, "rows"):
            blocks.append(table_block([[cell.text for cell in row.cells] for row in item.rows]))
        else:
            if item.text.strip():
                blocks.append(
                    {
                        "kind": "heading"
                        if item.style and item.style.name.startswith("Heading")
                        else "paragraph",
                        "text": item.text,
                    }
                )
            for equation in item._p.xpath(".//m:oMath"):
                blocks.append(
                    {
                        "kind": "equation",
                        "text": "".join(equation.itertext()),
                        "source_xml": etree.tostring(equation, encoding="unicode"),
                        "review": "Office Math retained; equation semantics and layout require review.",
                    }
                )
            for drawing in item._p.xpath(".//a:blip"):
                relation = drawing.get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                )
                if relation in document.part.related_parts:
                    blocks.append(
                        image_block(
                            document.part.related_parts[relation].blob, "Embedded document figure"
                        )
                    )
    return blocks, warnings


def extract_pdf(data: bytes) -> tuple[list[dict], list[str]]:
    """Extract bounded text, detected ruled tables, and raster figures; flag ambiguous structure."""
    if not data.startswith(b"%PDF-"):
        raise ValueError("Invalid PDF signature")
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted or len(reader.pages) > 30:
        raise ValueError("Encrypted PDF or more than 30 pages")
    blocks: list[dict] = []
    warnings = [
        "PDF reading order, graphs, equations, and detected table boundaries require human review."
    ]
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        for number, page in enumerate(reader.pages):
            content = page.get_contents()
            if content is not None and len(content.get_data()) > 4 * 1024 * 1024:
                raise ValueError("PDF page content stream exceeds limit")
            text = page.extract_text() or ""
            blocks.append({"kind": "heading", "text": f"Page {number + 1}"})
            blocks.append({"kind": "paragraph", "text": text})
            if not text.strip():
                warnings.append(f"Page {number + 1}: no text layer; OCR/vision adapter needed.")
            for rows in pdf.pages[number].extract_tables():
                blocks.append(table_block(rows))
            for figure in list(page.images)[:12]:
                blocks.append(image_block(figure.data, f"PDF figure, page {number + 1}"))
    return blocks, warnings


def extract(name: str, data: bytes) -> tuple[list[dict], list[str]]:
    """Select a supported importer by explicit extension and validate the real payload."""
    extension = name.rsplit(".", 1)[-1].lower()
    if extension not in SUPPORTED:
        raise ValueError("Unsupported format; register a trusted importer adapter")
    if not data or len(data) > MAX_BYTES:
        raise ValueError("File must contain 1 byte to 4 MiB")
    if extension == "docx":
        return extract_docx(data)
    if extension == "pdf":
        return extract_pdf(data)
    if extension in {"png", "jpg", "jpeg", "webp"}:
        return [image_block(data, name)], [
            "Raster retained; OCR/vision interpretation requires a configured adapter and review."
        ]
    if extension == "xlsx":
        archive_guard(data, "xl/workbook.xml")
        workbook = load_workbook(
            io.BytesIO(data), read_only=True, data_only=False, keep_links=False
        )
        blocks = []
        try:
            if len(workbook.worksheets) > 12:
                raise ValueError("Workbook exceeds sheet limit")
            for sheet in workbook:
                if sheet.max_row > 2000 or sheet.max_column > 40:
                    raise ValueError("Worksheet exceeds table limit")
                blocks.extend(
                    [{"kind": "heading", "text": sheet.title}, table_block(list(sheet.values))]
                )
        finally:
            workbook.close()
        return blocks, [
            "Workbook formulas retained as text, not evaluated; charts, images and macros are not interpreted."
        ]
    text = data.decode("utf-8-sig")
    if "\x00" in text or len(text) > MAX_TEXT:
        raise ValueError("Invalid text encoding or text exceeds character limit")
    if extension in {"csv", "tsv"}:
        return [
            table_block(
                list(csv.reader(io.StringIO(text), delimiter="\t" if extension == "tsv" else ","))
            )
        ], []
    if extension == "json":
        value = json.loads(
            text, parse_constant=lambda _: (_ for _ in ()).throw(ValueError("Non-finite JSON"))
        )
        if isinstance(value, list) and value and all(isinstance(row, dict) for row in value):
            keys = list(dict.fromkeys(key for row in value for key in row))
            return [table_block([keys] + [[row.get(key, "") for key in keys] for row in value])], []
        return [{"kind": "paragraph", "text": json.dumps(value, ensure_ascii=False, indent=2)}], []
    if extension == "xml":
        root = ElementTree.fromstring(text)
        return [{"kind": "paragraph", "text": " ".join(root.itertext())}], [
            "XML element tree retained as source; text-only canonical representation."
        ]
    if extension in {"html", "htm", "md"}:
        safe = clean_html(
            MarkdownIt("commonmark", {"html": False}).render(text) if extension == "md" else text
        )
        return [{"kind": "html", "text": safe}], [
            "Active content, remote assets, styles and unsafe links removed."
        ]
    return [{"kind": "paragraph", "text": part} for part in text.split("\n\n") if part.strip()], []


def render_document(
    title: str,
    blocks: list[dict],
    measurements: list[dict],
    provenance: dict,
    config: WorkspaceConfig,
    warnings: list[str],
) -> str:
    """Render all exports' canonical first stage using escaped text and trusted structure."""
    parts = [f"<h1>{escape(title)}</h1>"]
    for section in config.sections:
        parts.append(f"<section><h2>{section.title()}</h2>")
        if section == "overview":
            parts.append(
                "<p>Standardized document. Review source fidelity before operational use.</p>"
            )
            parts.extend(f"<p>{escape(warning)}</p>" for warning in warnings)
        elif section == "content":
            for block in blocks:
                kind = block["kind"]
                if kind == "table":
                    rows = block["rows"]
                    parts.append("<table><caption>Imported table</caption>")
                    for index, row in enumerate(rows):
                        tag = "th" if index == 0 else "td"
                        parts.append(
                            "<tr>"
                            + "".join(f"<{tag}>{escape(cell)}</{tag}>" for cell in row)
                            + "</tr>"
                        )
                    parts.append("</table>")
                elif kind == "figure":
                    parts.append(
                        f'<figure><img alt="{escape(block["text"], quote=True)}" src="data:image/png;base64,{block["image"]}"><figcaption>{escape(block["review"])}</figcaption></figure>'
                    )
                elif kind == "html":
                    parts.append(block["text"])
                else:
                    tag = "h3" if kind == "heading" else "pre" if kind == "equation" else "p"
                    parts.append(f"<{tag}>{escape(block['text'])}</{tag}>")
                    if block.get("review"):
                        parts.append(f"<p>{escape(block['review'])}</p>")
        elif section == "measurements":
            parts.append(
                "<ul>"
                + "".join(
                    f"<li>{escape(record['original'])} → {record['value']} {record['unit']}</li>"
                    for record in measurements
                )
                + "</ul>"
            )
        else:
            parts.append(f"<pre>{escape(json.dumps(provenance, indent=2))}</pre>")
        parts.append("</section>")
    style = "body{font:16px system-ui;max-width:900px;margin:auto;padding:24px;color:#18312c}table{border-collapse:collapse;width:100%}td,th{border:1px solid #bbb;padding:8px}img{max-width:100%;height:auto}pre{white-space:pre-wrap;overflow-wrap:anywhere}"
    return f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{escape(title)}</title><style>{style}</style></head><body>{"".join(parts)}</body></html>'


def ingest(name: str, data: bytes, config: WorkspaceConfig) -> dict:
    """Produce normalized semantic blocks, canonical HTML, and auditable conversion metadata."""
    blocks, warnings = extract(name, data)
    if len(blocks) > 3000 or sum(len(json.dumps(block)) for block in blocks) > 5 * 1024 * 1024:
        raise ValueError("Extracted content exceeds budget")
    measurements: list[dict] = []
    for index, block in enumerate(blocks):
        if block["kind"] == "table":
            for row in block["rows"]:
                for column, cell in enumerate(row):
                    row[column], records = normalize(cell, config.units)
                    measurements.extend({**record, "block": index} for record in records)
        elif block["kind"] not in {"figure", "equation"}:
            if block["kind"] == "html":
                soup = BeautifulSoup(block["text"], "html.parser")
                for node in list(soup.find_all(string=True)):
                    normalized, records = normalize(str(node), config.units)
                    node.replace_with(normalized)
                    measurements.extend({**record, "block": index} for record in records)
                block["text"] = str(soup)
            else:
                block["text"], records = normalize(block["text"], config.units)
                measurements.extend({**record, "block": index} for record in records)
    provenance = {
        "source_name": name,
        "sha256": hashlib.sha256(data).hexdigest(),
        "config_revision": config.revision,
        "conversion_path": [name.rsplit(".", 1)[-1].lower(), "html"],
        "schema_version": "1.0",
    }
    html = render_document(name, blocks, measurements, provenance, config, warnings)
    text = BeautifulSoup(html, "html.parser").get_text(" ", strip=True)
    return {
        "schema_version": "1.0",
        "title": name,
        "blocks": blocks,
        "measurements": measurements,
        "warnings": warnings,
        "provenance": provenance,
        "html": html,
        "text": text[:MAX_TEXT],
        "review_status": "needs-review"
        if warnings or any(block.get("review") for block in blocks)
        else "extracted",
    }
